import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


@dataclass
class NormalizedRecord:
    record_id: str
    data: dict[str, Any]
    location: dict[str, Any]


@dataclass
class NormalizedDataset:
    format: str
    record_count: int
    columns: list[str]
    records: list[NormalizedRecord]
    metadata: dict[str, Any]

    def to_canonical_dict(self) -> dict[str, Any]:
        """
        Produces a canonical dictionary suitable for deterministic serialization.
        """
        return {
            "format": self.format,
            "record_count": self.record_count,
            "columns": sorted(self.columns),
            "records": [
                {
                    "record_id": r.record_id,
                    "data": {k: str(v) if v is not None else "" for k, v in sorted(r.data.items())},
                    "location": {k: v for k, v in sorted(r.location.items())},
                }
                for r in self.records
            ],
        }


def validate_and_process(file_path: str | Path, extension: str) -> dict[str, Any]:
    """
    Validates dataset format and produces high-level summary.
    Maintains backward compatibility with earlier endpoint responses.
    """
    normalized = normalize_dataset(file_path, extension)
    summary: dict[str, Any] = {
        "format": normalized.format,
        "record_count": normalized.record_count,
        "columns": normalized.columns,
        "column_count": len(normalized.columns),
    }
    summary.update(normalized.metadata)
    return summary


def normalize_dataset(file_path: str | Path, extension: str) -> NormalizedDataset:
    path = Path(file_path)
    ext = extension.lower().strip()
    if not ext.startswith("."):
        ext = f".{ext}"

    if ext == ".csv":
        return _normalize_csv(path)
    elif ext == ".json":
        return _normalize_json(path)
    elif ext == ".txt":
        return _normalize_txt(path)
    elif ext == ".xlsx":
        return _normalize_xlsx(path)
    else:
        raise ValueError(f"Unsupported dataset format: {extension}")


def _normalize_csv(path: Path) -> NormalizedDataset:
    df = pd.read_csv(path)
    # Fill NaN with empty string
    df = df.fillna("")
    columns = [str(c) for c in df.columns]
    records: list[NormalizedRecord] = []

    for idx, row in df.iterrows():
        rec_id = str(idx)
        data = {col: str(row[col]) for col in columns}
        location = {
            "row": int(idx) + 1,
            "column": "all",
            "sheet": "default",
        }
        records.append(NormalizedRecord(record_id=rec_id, data=data, location=location))

    return NormalizedDataset(
        format="csv",
        record_count=len(records),
        columns=columns,
        records=records,
        metadata={"rows": len(df), "columns": columns},
    )


def _normalize_json(path: Path) -> NormalizedDataset:
    with path.open("r", encoding="utf-8") as f:
        raw_data = json.load(f)

    records: list[NormalizedRecord] = []
    columns_set: set[str] = set()

    if isinstance(raw_data, list):
        for idx, item in enumerate(raw_data):
            rec_id = str(idx)
            if isinstance(item, dict):
                data = {str(k): str(v) for k, v in item.items()}
                columns_set.update(data.keys())
            else:
                data = {"value": str(item)}
                columns_set.add("value")

            location = {"record_index": idx, "row": idx + 1, "column": "all"}
            records.append(NormalizedRecord(record_id=rec_id, data=data, location=location))

        columns = sorted(list(columns_set))
        return NormalizedDataset(
            format="json",
            record_count=len(records),
            columns=columns,
            records=records,
            metadata={"top_level_type": "array", "records": len(records)},
        )

    elif isinstance(raw_data, dict):
        for key, val in raw_data.items():
            if isinstance(val, dict):
                data = {str(k): str(v) for k, v in val.items()}
                columns_set.update(data.keys())
            else:
                data = {"content": str(val)}
                columns_set.add("content")

            location = {"key": str(key), "column": "all"}
            records.append(NormalizedRecord(record_id=str(key), data=data, location=location))

        columns = sorted(list(columns_set))
        return NormalizedDataset(
            format="json",
            record_count=len(records),
            columns=columns,
            records=records,
            metadata={"top_level_type": "object", "records": len(records)},
        )
    else:
        raise ValueError("JSON file must contain an array or object at top level.")


def _normalize_txt(path: Path) -> NormalizedDataset:
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    records: list[NormalizedRecord] = []

    for line_num, line in enumerate(lines, start=1):
        rec_id = str(line_num)
        data = {"text": line}
        location = {
            "line_number": line_num,
            "row": line_num,
            "column": "text",
        }
        records.append(NormalizedRecord(record_id=rec_id, data=data, location=location))

    return NormalizedDataset(
        format="txt",
        record_count=len(records),
        columns=["text"],
        records=records,
        metadata={"characters": len(text), "lines": len(lines)},
    )


def _normalize_xlsx(path: Path) -> NormalizedDataset:
    workbook = openpyxl.load_workbook(path, data_only=True)
    records: list[NormalizedRecord] = []
    columns_set: set[str] = set()

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        header = [str(c or f"Col_{i}") for i, c in enumerate(rows[0])]
        columns_set.update(header)

        for row_idx, row_values in enumerate(rows[1:], start=2):
            rec_id = f"{sheet_name}_r{row_idx}"
            data = {}
            for col_idx, col_name in enumerate(header):
                val = row_values[col_idx] if col_idx < len(row_values) else ""
                data[col_name] = str(val) if val is not None else ""

            location = {
                "sheet": sheet_name,
                "row": row_idx,
                "column": "all",
            }
            records.append(NormalizedRecord(record_id=rec_id, data=data, location=location))

    columns = sorted(list(columns_set))
    return NormalizedDataset(
        format="xlsx",
        record_count=len(records),
        columns=columns,
        records=records,
        metadata={"sheets": workbook.sheetnames, "rows": len(records)},
    )


def write_modified_dataset(
    original_path: Path,
    output_path: Path,
    file_type: str,
    modified_records: list[dict[str, Any]],
    removed_record_ids: set[str],
) -> None:
    """
    Saves a modified version of the dataset for Version 2 creation,
    applying cell edits and row removals while preserving format.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fmt = file_type.lower().lstrip(".")

    # Convert modified_records list to lookup by record_id
    mod_lookup = {str(r["record_id"]): r["data"] for r in modified_records if "record_id" in r and "data" in r}

    normalized = normalize_dataset(original_path, f".{fmt}")

    kept_records: list[NormalizedRecord] = []
    for r in normalized.records:
        if r.record_id in removed_record_ids:
            continue
        # Apply edits if any
        if r.record_id in mod_lookup:
            updated_data = dict(r.data)
            updated_data.update(mod_lookup[r.record_id])
            kept_records.append(NormalizedRecord(record_id=r.record_id, data=updated_data, location=r.location))
        else:
            kept_records.append(r)

    if fmt == "csv":
        rows = [r.data for r in kept_records]
        df = pd.DataFrame(rows, columns=normalized.columns)
        df.to_csv(output_path, index=False)

    elif fmt == "json":
        if normalized.metadata.get("top_level_type") == "object":
            obj_data = {r.record_id: r.data for r in kept_records}
            with output_path.open("w", encoding="utf-8") as f:
                json.dump(obj_data, f, indent=2)
        else:
            list_data = [r.data for r in kept_records]
            with output_path.open("w", encoding="utf-8") as f:
                json.dump(list_data, f, indent=2)

    elif fmt == "txt":
        lines = [r.data.get("text", "") for r in kept_records]
        output_path.write_text("\n".join(lines), encoding="utf-8")

    elif fmt == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(normalized.columns)
        for r in kept_records:
            row_vals = [r.data.get(col, "") for col in normalized.columns]
            ws.append(row_vals)
        wb.save(output_path)
    else:
        raise ValueError(f"Cannot serialize format: {fmt}")
